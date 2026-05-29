from __future__ import annotations
import io
from dataclasses import dataclass, field
from typing import Optional, Union

import openpyxl

PUZZLE_CHAIN_NAMES = {"Puzzle Chain 1", "Puzzle Chain 2", "Puzzle Chain 3", "Puzzle Chain 4"}


@dataclass
class LootEntry:
    item: str
    probability: float


@dataclass
class LootTable:
    name: str
    default_item: str
    default_probability: float
    entries: list[LootEntry] = field(default_factory=list)


@dataclass
class Plant:
    prefab: str
    harvest_time: float
    max_harvests: int
    loot_table_name: str  # loot table name OR direct item prefab
    on_die: Optional[str] = None  # item dropped when all charges exhausted


@dataclass
class MergeChain:
    name: str
    items: list[str]  # ordered low to high level (index 0 = base)


@dataclass
class ZoneData:
    zone_id: int
    zone_type: str
    tile_count: int
    composition: dict[str, float]  # prefab -> % (0–100)


@dataclass
class ZoneUnlock:
    zone_id: int
    required_currency_level: Optional[int]  # None = no unlock required


@dataclass
class GameData:
    zones: dict[int, ZoneData]
    plants: dict[str, Plant]
    loot_tables: dict[str, LootTable]
    chains: list[MergeChain]
    currency_chain: MergeChain
    zone_unlocks: dict[int, ZoneUnlock]
    points_chain: MergeChain
    point_values: dict[str, int]


def _open_wb(src: Union[str, bytes, io.BytesIO]) -> openpyxl.Workbook:
    if isinstance(src, bytes):
        return openpyxl.load_workbook(io.BytesIO(src), data_only=True)
    if isinstance(src, io.BytesIO):
        src.seek(0)
        return openpyxl.load_workbook(src, data_only=True)
    return openpyxl.load_workbook(src, data_only=True)


def load_data(
    layout_path: Union[str, bytes, io.BytesIO],
    ge_revamp_path: Optional[Union[str, bytes, io.BytesIO]] = None,
) -> GameData:
    layout_wb = _open_wb(layout_path)
    ge_revamp_wb = _open_wb(ge_revamp_path) if ge_revamp_path is not None else None

    _ZONE_GATES_SHEET = next(
        (s for s in layout_wb.sheetnames if s in ("Zone Gates", "Zone Gating")), None
    )
    _ZONE_GATING_FORMAT = _ZONE_GATES_SHEET == "Zone Gating"
    is_new_layout = _ZONE_GATES_SHEET is not None

    zones = _parse_zones(layout_wb)
    loot_tables = _parse_loot_layout(layout_wb)
    plants = _parse_plants_layout(layout_wb, is_new_layout, loot_tables)
    chains, currency_chain = _parse_chains(layout_wb, is_new_layout)
    points_chain, point_values = _parse_points_layout(layout_wb, is_new_layout)

    if ge_revamp_wb:
        loot_tables.update(_parse_loot_event(ge_revamp_wb))
        ge_overrides = _parse_ge_revamp_objects(ge_revamp_wb)
        for prefab in list(plants.keys()):
            override = ge_overrides.get(prefab)
            if override:
                charges, seconds, on_die, harvest_loot = override
                p = plants[prefab]
                plants[prefab] = Plant(
                    prefab=p.prefab,
                    harvest_time=seconds,
                    max_harvests=charges,
                    loot_table_name=harvest_loot if harvest_loot else p.loot_table_name,
                    on_die=on_die,
                )
        for prefab, (charges, seconds, on_die, harvest_loot) in ge_overrides.items():
            if prefab not in plants and harvest_loot:
                plants[prefab] = Plant(
                    prefab=prefab,
                    harvest_time=seconds,
                    max_harvests=charges,
                    loot_table_name=harvest_loot,
                    on_die=on_die,
                )

    if is_new_layout and _ZONE_GATING_FORMAT:
        zone_unlocks = _parse_zone_unlocks_gating(layout_wb)
    elif is_new_layout:
        zone_unlocks = _parse_zone_unlocks_gates(layout_wb, _ZONE_GATES_SHEET)
    else:
        zone_unlocks = _parse_zone_unlocks_old(layout_wb)

    max_chain_level = len(currency_chain.items)
    for zone_id in list(zone_unlocks.keys()):
        if zone_id >= 3:
            level = min(zone_id, max_chain_level)
            zone_unlocks[zone_id] = ZoneUnlock(zone_id=zone_id, required_currency_level=level)

    return GameData(
        zones=zones,
        plants=plants,
        loot_tables=loot_tables,
        chains=chains,
        currency_chain=currency_chain,
        zone_unlocks=zone_unlocks,
        points_chain=points_chain,
        point_values=point_values,
    )


# ── Zone parsing ──────────────────────────────────────────────────────────────

_BALANCE_SHEET_NAMES = [
    "Level_Event_LakeCottage_Balance",
    "Copy of Level_Event_LakeCottage",
]


def _parse_zones(wb) -> dict[int, ZoneData]:
    sheet_name = next((s for s in _BALANCE_SHEET_NAMES if s in wb.sheetnames), None)
    if sheet_name is None:
        raise ValueError(f"No zone layout sheet found. Expected one of: {_BALANCE_SHEET_NAMES}")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))

    zone_types = rows[0]
    zone_nums = rows[1]
    tile_counts = rows[2]

    zones: dict[int, ZoneData] = {}
    max_col = len(zone_nums)

    for col in range(1, max_col):
        zn = zone_nums[col]
        zt = zone_types[col]
        tc = tile_counts[col]
        if zn is None or zt is None or tc is None:
            continue
        try:
            zone_id = int(zn)
            tile_count = int(tc)
        except (ValueError, TypeError):
            continue

        composition: dict[str, float] = {}
        row_idx = 5
        while row_idx + 1 < len(rows):
            row_a = rows[row_idx]
            label_a = row_a[0] if row_a else None
            if label_a == "Prefab":
                row_b = rows[row_idx + 1]
                prefab_val = row_a[col] if col < len(row_a) else None
                pct_val = row_b[col] if col < len(row_b) else None
                if prefab_val and pct_val and str(prefab_val) not in ("Ignore",):
                    try:
                        composition[str(prefab_val)] = float(pct_val)
                    except (ValueError, TypeError):
                        pass
            row_idx += 1

        zones[zone_id] = ZoneData(
            zone_id=zone_id,
            zone_type=str(zt),
            tile_count=tile_count,
            composition=composition,
        )

    return zones


# ── Loot table parsing ────────────────────────────────────────────────────────

def _parse_loot_layout(wb) -> dict[str, LootTable]:
    """Parse Generator LootTable sheet (same column layout for old and new layout file)."""
    loot_ws = wb["Generator LootTable"]
    loot_tables: dict[str, LootTable] = {}

    for row in loot_ws.iter_rows(min_row=2, max_row=loot_ws.max_row, values_only=True):
        if not row[0]:
            continue
        lt_name = str(row[1]) if row[1] else str(row[0])
        default_item = str(row[2]) if row[2] else ""
        default_prob = float(row[4]) if row[4] else 0.0

        entries: list[LootEntry] = []
        for i in range(7):
            item_col = 5 + i * 4
            input_col = 8 + i * 4
            if item_col < len(row) and input_col < len(row):
                item_val = row[item_col]
                prob_val = row[input_col]
                if item_val and prob_val:
                    entries.append(LootEntry(item=str(item_val), probability=float(prob_val)))

        loot_tables[lt_name] = LootTable(
            name=lt_name,
            default_item=default_item,
            default_probability=default_prob,
            entries=entries,
        )

    return loot_tables


def _parse_loot_event(loot_wb) -> dict[str, LootTable]:
    """Parse Loot_Tables sheet from the Loot file."""
    loot_tables: dict[str, LootTable] = {}
    sheet_name = next(
        (s for s in loot_wb.sheetnames if "Loot_Tables" in s or "Loot Tables" in s),
        None,
    )
    if not sheet_name:
        return loot_tables

    ws = loot_wb[sheet_name]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[0]:
            continue
        lt_name = str(row[0])
        default_item = str(row[1]) if row[1] else ""
        default_prob = float(row[3]) if row[3] else 0.0

        entries: list[LootEntry] = []
        for i in range(7):
            item_col = 4 + i * 4
            prob_col = 7 + i * 4
            if item_col < len(row):
                item_val = row[item_col]
                prob_val = row[prob_col] if prob_col < len(row) else None
                if item_val and prob_val:
                    entries.append(LootEntry(item=str(item_val), probability=float(prob_val)))

        loot_tables[lt_name] = LootTable(
            name=lt_name,
            default_item=default_item,
            default_probability=default_prob,
            entries=entries,
        )

    return loot_tables


# ── Plant parsing ─────────────────────────────────────────────────────────────

def _parse_ge_revamp_objects(ge_revamp_wb) -> dict[str, tuple[int, float, Optional[str], Optional[str]]]:
    """Return prefab → (charges, harvest_seconds, on_die, harvest_loot) from GE Revamp Data Editor.

    Also registers entries under stripped name (Brambles_GERevamp → Brambles)
    for compatibility with current layout prefab names.
    """
    ws = ge_revamp_wb["Object_Definitions - GERevamp"]
    result: dict[str, tuple[int, float, Optional[str], Optional[str]]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) <= 102:
            continue
        prefab = row[3]
        charges = row[99]
        seconds = row[97]
        on_die = row[102]
        harvest_loot = row[95] if len(row) > 95 else None
        if not prefab or charges is None:
            continue
        entry = (int(charges), float(seconds) if seconds else 0.0, str(on_die) if on_die else None, str(harvest_loot) if harvest_loot else None)
        result[str(prefab)] = entry
        stripped = str(prefab).replace("_GERevamp", "")
        if stripped != str(prefab):
            result[stripped] = entry
    return result


def _parse_plants_layout(wb, is_new: bool, loot_tables: dict[str, LootTable]) -> dict[str, Plant]:
    """Parse harvestable items from Item Database. Column offsets differ between old/new layout."""
    item_ws = wb["Item Database"]
    # old: Chain=0, Prefab=3, Harvestable=5, HarvestTime=7, Loottable=8, HarvestCount=9
    # new: Chain=0, ChainLevel=1, Prefab=4, Harvestable=6, HarvestTime=8, Loottable=9, HarvestCount=10
    if is_new:
        COL_PREFAB, COL_HARVEST, COL_TIME, COL_LOOT, COL_COUNT = 4, 6, 8, 9, 10
    else:
        COL_PREFAB, COL_HARVEST, COL_TIME, COL_LOOT, COL_COUNT = 3, 5, 7, 8, 9

    plants: dict[str, Plant] = {}
    for row in item_ws.iter_rows(min_row=2, max_row=item_ws.max_row, values_only=True):
        if len(row) <= COL_COUNT:
            continue
        prefab = row[COL_PREFAB]
        harvestable = row[COL_HARVEST]
        harvest_time = row[COL_TIME]
        loot_name = row[COL_LOOT]
        harvest_count = row[COL_COUNT]

        if prefab and harvestable == "x" and harvest_time and harvest_count:
            plants[str(prefab)] = Plant(
                prefab=str(prefab),
                harvest_time=float(harvest_time),
                max_harvests=int(harvest_count),
                loot_table_name=str(loot_name) if loot_name else "",
            )

    # Wire plants to loot table names from Generator LootTable
    loot_ws = wb["Generator LootTable"]
    for row in loot_ws.iter_rows(min_row=2, max_row=loot_ws.max_row, values_only=True):
        if not row[0]:
            continue
        main_prefab = str(row[0])
        lt_name = str(row[1]) if row[1] else main_prefab
        if main_prefab in plants:
            old = plants[main_prefab]
            plants[main_prefab] = Plant(
                prefab=old.prefab,
                harvest_time=old.harvest_time,
                max_harvests=old.max_harvests,
                loot_table_name=lt_name,
            )

    return plants


def _parse_plants_objects(objects_wb) -> dict[str, Plant]:
    """Parse harvestable items from Object_Definitions sheet."""
    ws = objects_wb["Object_Definitions - LakeCott"]
    # col5=Prefab, col104=HarvestLoot, col106=HarvestSeconds, col108=HarvestCharges
    plants: dict[str, Plant] = {}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if len(row) <= 108:
            continue
        prefab = row[5]
        harvest_loot = row[104]
        harvest_sec = row[106]
        harvest_charges = row[108]

        if prefab and harvest_loot and harvest_sec is not None:
            plants[str(prefab)] = Plant(
                prefab=str(prefab),
                harvest_time=float(harvest_sec),
                max_harvests=int(harvest_charges) if harvest_charges else 1,
                loot_table_name=str(harvest_loot),
            )

    return plants


# ── Chain parsing ─────────────────────────────────────────────────────────────

def _parse_chains(wb, is_new: bool) -> tuple[list[MergeChain], MergeChain]:
    item_ws = wb["Item Database"]
    # old: Chain=0, Prefab=3, Mergeable=4
    # new: Chain=0, ChainLevel=1, Prefab=4, Mergeable=5
    COL_PREFAB = 4 if is_new else 3
    COL_MERGE = 5 if is_new else 4

    chain_items: dict[str, list[str]] = {}
    currency_items: list[str] = []

    for row in item_ws.iter_rows(min_row=2, max_row=item_ws.max_row, values_only=True):
        if len(row) <= COL_MERGE:
            continue
        chain_name = row[0]
        prefab = row[COL_PREFAB]
        mergeable = row[COL_MERGE]

        if not prefab or not chain_name:
            continue

        chain_str = str(chain_name)
        if chain_str == "Discovery Chain":
            currency_items.append(str(prefab))
        elif mergeable == "x" and chain_str in PUZZLE_CHAIN_NAMES:
            if chain_str not in chain_items:
                chain_items[chain_str] = []
            chain_items[chain_str].append(str(prefab))

    chains = [MergeChain(name=n, items=items) for n, items in chain_items.items()]
    currency_chain = MergeChain(name="Discovery Chain", items=currency_items)
    return chains, currency_chain


# ── Points chain parsing ──────────────────────────────────────────────────────

def _parse_points_layout(wb, is_new: bool) -> tuple[MergeChain, dict[str, int]]:
    item_ws = wb["Item Database"]
    COL_PREFAB = 4 if is_new else 3
    # new format: col 11 = Misc = point value; old: approximate with 3^i
    COL_MISC = 11 if is_new else None

    point_items: list[str] = []
    point_values: dict[str, int] = {}

    for row in item_ws.iter_rows(min_row=2, max_row=item_ws.max_row, values_only=True):
        if len(row) <= COL_PREFAB:
            continue
        chain_name = row[0]
        prefab = row[COL_PREFAB]
        if chain_name and "Point" in str(chain_name) and prefab:
            point_items.append(str(prefab))
            if COL_MISC and COL_MISC < len(row) and row[COL_MISC] is not None:
                try:
                    point_values[str(prefab)] = int(float(row[COL_MISC]))
                except (ValueError, TypeError):
                    pass

    if not point_values:
        point_values = {p: 3 ** i for i, p in enumerate(point_items)}

    return MergeChain(name="Points Chain", items=point_items), point_values


def _parse_points_objects(objects_wb, layout_wb, is_new: bool) -> tuple[MergeChain, dict[str, int]]:
    """Use Strength (col42) from Object_Definitions as event point value."""
    ws = objects_wb["Object_Definitions - LakeCott"]
    point_values: dict[str, int] = {}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if len(row) <= 42:
            continue
        prefab = row[5]
        strength = row[42]
        if prefab and strength is not None:
            try:
                point_values[str(prefab)] = int(float(strength))
            except (ValueError, TypeError):
                pass

    # Points chain items (ordered) from layout Item Database
    points_chain, _ = _parse_points_layout(layout_wb, is_new)
    # Supplement point_values with Points Chain items if missing
    for i, prefab in enumerate(points_chain.items):
        if prefab not in point_values:
            point_values[prefab] = 3 ** i

    return points_chain, point_values


# ── Zone unlock parsing ───────────────────────────────────────────────────────

def _parse_zone_unlocks_old(wb) -> dict[int, ZoneUnlock]:
    if "Zone unlock" not in wb.sheetnames:
        raise KeyError(
            f"Layout file missing required sheet. Expected 'Zone Gates' or 'Zone Gating' (new format) or "
            f"'Zone unlock' (old format). Found sheets: {wb.sheetnames}"
        )
    ws = wb["Zone unlock"]
    unlocks: dict[int, ZoneUnlock] = {}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[1] is None:
            continue
        try:
            zone_id = int(row[1])
        except (ValueError, TypeError):
            continue
        item = row[2]
        if item:
            try:
                level = int(str(item).split("_")[-1])
            except ValueError:
                level = None
            unlocks[zone_id] = ZoneUnlock(zone_id=zone_id, required_currency_level=level)
        else:
            unlocks[zone_id] = ZoneUnlock(zone_id=zone_id, required_currency_level=None)

    return unlocks


def _parse_zone_unlocks_gates(wb, sheet_name: str = "Zone Gates") -> dict[int, ZoneUnlock]:
    """Parse Zone Gates/Gating matrix: for zone N, required level = highest Discovery Chain
    currency level whose cumulative count becomes > 0 at or before zone N."""
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    if not rows:
        return {}

    headers = rows[0]
    zone_cols: dict[int, int] = {}
    for col_idx, h in enumerate(headers):
        if h and str(h).startswith("Zone "):
            try:
                zone_id = int(str(h).split()[-1])
                zone_cols[zone_id] = col_idx
            except ValueError:
                pass

    # Collect Discovery Chain rows with their currency level
    discovery: list[tuple[int, tuple]] = []
    for row in rows[1:]:
        if row[0] != "Discovery Chain" or not row[1]:
            continue
        try:
            level = int(str(row[1]).split("_")[-1])
            discovery.append((level, row))
        except (ValueError, IndexError):
            pass

    unlocks: dict[int, ZoneUnlock] = {}
    for zone_id, col_idx in sorted(zone_cols.items()):
        max_level: Optional[int] = None
        for level, row in discovery:
            if col_idx < len(row) and row[col_idx] is not None:
                try:
                    if float(row[col_idx]) > 0:
                        if max_level is None or level > max_level:
                            max_level = level
                except (ValueError, TypeError):
                    pass
        unlocks[zone_id] = ZoneUnlock(zone_id=zone_id, required_currency_level=max_level)

    return unlocks


def _parse_zone_unlocks_gating(wb) -> dict[int, ZoneUnlock]:
    """Parse Zone Gating sheet: rows are (zone_name_or_NA, fog_prefab, currency_item).
    Zone N unlocks when the player has the currency item listed in col 2."""
    ws = wb["Zone Gating"]
    unlocks: dict[int, ZoneUnlock] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row[0] or not row[2]:
            continue
        zone_str = str(row[0]).strip()
        if zone_str == "N/A":
            continue
        try:
            zone_id = int(zone_str.split()[-1])
        except (ValueError, IndexError):
            continue
        currency = str(row[2]).strip()
        try:
            level = int(currency.split("_")[-1])
        except (ValueError, IndexError):
            level = None
        unlocks[zone_id] = ZoneUnlock(zone_id=zone_id, required_currency_level=level)
    return unlocks
